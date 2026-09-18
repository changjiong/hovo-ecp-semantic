#!/usr/bin/env python3
"""Split declared PDF and DOCX sources into auditable source units."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ARTICLE_RE = re.compile(r"(?m)^\s*(第[一二三四五六七八九十百零〇]+条)[ \t]+")
SECTION_RE = re.compile(r"(?m)^\s*((?:\d+\.)+\d+)\s+")
CLAUSE_RE = re.compile(r"(?m)^[ \t]*([（(][一二三四五六七八九十百零〇]+[）)])[ \t]*")


def digest(text: str) -> str:
    return "sha256:" + hashlib.sha256(text.encode("utf-8")).hexdigest()


def normalized(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return "\n".join(line.rstrip() for line in text.replace("\r\n", "\n").splitlines()).strip()


def split_block(text: str, page: int | None = None) -> list[tuple[str, str, str]]:
    text = normalized(text)
    if not text:
        label = f"第{page}页" if page is not None else "空白内容"
        return [("PAGE_BLOCK", label, "[UNREADABLE_OR_EMPTY]")]
    matches = list(ARTICLE_RE.finditer(text)) or list(SECTION_RE.finditer(text))
    if not matches:
        label = f"第{page}页正文" if page is not None else "正文段落"
        return split_clauses("PAGE_BLOCK", label, text)
    units: list[tuple[str, str, str]] = []
    if matches[0].start() > 0:
        prefix = normalized(text[: matches[0].start()])
        if prefix:
            units.extend(split_clauses("PREAMBLE" if page in (None, 1) else "PAGE_BLOCK", "前置或续接内容", prefix))
    kind = "ARTICLE" if ARTICLE_RE.match(matches[0].group(0)) else "SECTION"
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        units.extend(split_clauses(kind, match.group(1), normalized(text[match.start() : end])))
    return units


def split_clauses(kind: str, label: str, text: str) -> list[tuple[str, str, str]]:
    matches = list(CLAUSE_RE.finditer(text))
    if not matches:
        return [(kind, label, text)]
    units: list[tuple[str, str, str]] = []
    prefix = normalized(text[: matches[0].start()])
    if prefix:
        units.append((kind, label, prefix))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        clause_label = f"{label}{match.group(1)}" if kind == "ARTICLE" else f"{label} {match.group(1)}"
        units.append(("CLAUSE", clause_label, normalized(text[match.start() : end])))
    return units


def pdf_units(path: Path) -> tuple[list[tuple[str, str, str, str]], str, str]:
    reader = PdfReader(path)
    units: list[tuple[str, str, str, str]] = []
    unreadable = []
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if not text.strip():
            unreadable.append(page_number)
        for kind, label, body in split_block(text, page_number):
            units.append((kind, label, f"PDF第{page_number}页 {label}", body))
    status = "PARTIAL" if unreadable else "COMPLETE"
    limitation = "无法提取文本的页: " + ", ".join(map(str, unreadable)) if unreadable else "未发现无法提取文本的页面；版式和OCR准确性仍需人工复核。"
    return units, status, limitation


def docx_units(path: Path) -> tuple[list[tuple[str, str, str, str]], str, str]:
    document = Document(path)
    units: list[tuple[str, str, str, str]] = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, 1):
        text = normalized(paragraph.text)
        if not text:
            continue
        for kind, label, body in split_block(text):
            units.append((kind, label, f"DOCX段落{paragraph_number} {label}", body))
    for table_number, table in enumerate(document.tables, 1):
        for row_number, row in enumerate(table.rows, 1):
            body = normalized("\t".join(normalized(cell.text) for cell in row.cells))
            if body:
                label = f"表格{table_number}第{row_number}行"
                units.append(("TABLE", label, f"DOCX{label}", body))
    status = "PARTIAL" if units else "FAILED"
    limitation = "段落和表格已提取；文本框、批注、页眉页脚和嵌入对象未覆盖，需人工复核。" if units else "未提取到可读段落或表格。"
    return units, status, limitation


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="domain-knowledge input.json")
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True, help="写入包含 source_units 的新 input.json")
    args = parser.parse_args()
    project_root = args.project_root.resolve()
    payload = json.loads(args.input.read_text(encoding="utf-8"))
    source_units = []
    extractions = []
    for source in payload["sources"]:
        source_id = source["source_id"]
        source_path = (project_root / source["artifact"]["path"]).resolve()
        if project_root not in source_path.parents or any("trash" in part.casefold() for part in source_path.parts):
            raise ValueError(f"非法来源路径: {source_path}")
        suffix = source_path.suffix.casefold()
        if suffix == ".pdf":
            rows, status, limitations = pdf_units(source_path)
            method = "pypdf page extraction with article/section segmentation"
        elif suffix == ".docx":
            rows, status, limitations = docx_units(source_path)
            method = "python-docx paragraph and table extraction"
        else:
            raise ValueError(f"不支持的来源格式: {source_path.name}")
        unit_ids = []
        for number, (kind, label, locator, text) in enumerate(rows, 1):
            unit_id = f"{source_id}.U{number:04d}"
            unit_ids.append(unit_id)
            source_units.append({
                "unit_id": unit_id,
                "source_id": source_id,
                "kind": kind,
                "label": label,
                "locator": locator,
                "text": text,
                "digest": digest(text),
            })
        extractions.append({
            "source_id": source_id,
            "status": status,
            "method": method,
            "unit_ids": unit_ids,
            "limitations": limitations,
        })
    payload["source_units"] = source_units
    payload["source_extractions"] = extractions
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
